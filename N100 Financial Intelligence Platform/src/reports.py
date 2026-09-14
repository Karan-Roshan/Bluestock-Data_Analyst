
"""Sprint 3 Excel deliverables."""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill,Font,Alignment
from openpyxl.utils import get_column_letter
from screener.engine import run_all,load_config
from analytics.peer import compute_percentiles,_peers,METRICS

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/"output"; OUT.mkdir(exist_ok=True)
GREEN=PatternFill("solid",fgColor="C6EFCE"); RED=PatternFill("solid",fgColor="FFC7CE")
YELLOW=PatternFill("solid",fgColor="FFEB9C"); GOLD=PatternFill("solid",fgColor="FFD966")

def write_screener():
    results=run_all(); cfg=load_config()["presets"]; path=OUT/"screener_output.xlsx"
    kpis=["roe","roce","npm","de","fcf","revenue_cagr_5yr","pat_cagr_5yr","eps_cagr_5yr","opm",
          "pe_ratio","pb_ratio","dividend_yield_pct","dividend_payout","icr","market_cap",
          "net_profit","asset_turnover","sales","composite_quality_score","sector_relative_score"]
    with pd.ExcelWriter(path,engine="openpyxl") as w:
        for name,df in results.items():
            cols=["company_id","company_name","broad_sector"]+[c for c in kpis if c in df.columns]
            df[cols].to_excel(w,sheet_name=name[:31],index=False)
    wb=load_workbook(path)
    for name,ws in zip(results,wb.worksheets):
        rules=cfg[name]; headers={c.value:i+1 for i,c in enumerate(ws[1])}
        for k,t in rules.items():
            if k.endswith("_min"):
                col=k[:-4]
                if col=="dividend_yield": col="dividend_yield_pct"
                if col in headers:
                    for r in range(2,ws.max_row+1):
                        v=ws.cell(r,headers[col]).value; ws.cell(r,headers[col]).fill=GREEN if v is not None and v>=t else RED
            elif k.endswith("_max"):
                col=k[:-4]
                if col=="dividend_payout": col="dividend_payout"
                if col=="debt_to_equity": col="de"
                if col in headers:
                    for r in range(2,ws.max_row+1):
                        v=ws.cell(r,headers[col]).value; ws.cell(r,headers[col]).fill=GREEN if v is not None and v<=t else RED
        for c in ws[1]: c.font=Font(bold=True); c.alignment=Alignment(horizontal="center")
        ws.freeze_panes="A2"
        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(12,max(len(str(x.value or "")) for x in col)+2),28)
    wb.save(path); return path

def write_peer_comparison():
    pct=compute_percentiles(); peers=_peers(); path=OUT/"peer_comparison.xlsx"
    base,_p,_b=__import__("screener.engine",fromlist=["load_universe"]).load_universe()
    groups=sorted(peers.group_name.dropna().astype(str).unique())
    with pd.ExcelWriter(path,engine="openpyxl") as w:
        for group in groups:
            ids=peers.loc[peers.group_name.astype(str)==group,"company_id"].drop_duplicates()
            b=base[base.company_id.isin(ids)].copy()
            val=pct[pct.peer_group_name==group].pivot_table(index="company_id",columns="metric",values="value")
            per=pct[pct.peer_group_name==group].pivot_table(index="company_id",columns="metric",values="percentile_rank")
            val.columns=[f"{x} Value" for x in val.columns]; per.columns=[f"{x} Percentile" for x in per.columns]
            cols=["company_id","company_name"]+[c for c in b.columns if c in [x[1] for x in METRICS.values()]]
            out=b[["company_id","company_name"]+[x for x in [m for m in METRICS.values()] if x in b]].merge(val,on="company_id",how="left").merge(per,on="company_id",how="left")
            out.to_excel(w,sheet_name=group[:31],index=False)
    wb=load_workbook(path)
    for group,ws in zip(groups,wb.worksheets):
        benchmark=set(peers.loc[(peers.group_name==group)&(peers.get("is_benchmark",False)==True),"company_id"]) if "is_benchmark" in peers else set()
        for r in range(2,ws.max_row+1):
            cid=ws.cell(r,1).value
            if cid in benchmark:
                for c in range(1,ws.max_column+1): ws.cell(r,c).fill=GOLD
            for c in range(1,ws.max_column+1):
                if "Percentile" in str(ws.cell(1,c).value):
                    v=ws.cell(r,c).value
                    if v is not None: ws.cell(r,c).fill=GREEN if v>=.75 else RED if v<=.25 else YELLOW
        for c in ws[1]: c.font=Font(bold=True)
        ws.freeze_panes="A2"
    wb.save(path); return path
